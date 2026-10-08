"""
Reads a folder of deal documents (deck .pptx, statements .pdf, model .xlsx)
and turns every number it can recognise into a Fact with its source location.

It is regex/label based, so it works on documents laid out like the sample
deal room and will miss things on messier files. Anything it can't read is
skipped quietly, which means "no finding" is not the same as "checked".
"""
import re
from dataclasses import dataclass
from pathlib import Path

ROLE_BY_EXT = {".pptx": "deck", ".pdf": "audited", ".xlsx": "model"}

LABELS = {
    "revenue": {"revenue", "revenues", "total revenue", "total revenues", "net sales", "net revenue", "sales"},
    "gross_profit": {"gross profit"},
    "gross_margin": {"gross margin", "gross margin %"},
    "ebitda": {"ebitda"},
    "ebitda_margin": {"ebitda margin", "ebitda margin %"},
    "operating_income": {"operating income", "ebit", "income from operations"},
    "da": {"depreciation and amortization", "depreciation & amortization", "d&a"},
    "net_income": {"net income", "net profit", "profit for the year"},
    "cash": {"cash", "cash and cash equivalents", "cash & equivalents"},
    "total_debt": {"total debt", "borrowings", "total borrowings", "debt"},
    "net_debt": {"net debt"},
    "capex": {"capital expenditures", "capital expenditure", "capex"},
    "customer_concentration": {"top customer", "largest customer"},
    "arr": {"arr", "annual recurring revenue"},
}
LABEL_TO_METRIC = {lab: m for m, labs in LABELS.items() for lab in labs}
PCT_METRICS = {"gross_margin", "ebitda_margin", "customer_concentration"}

SCALE = {"m": 1e6, "mm": 1e6, "million": 1e6, "b": 1e9, "bn": 1e9, "billion": 1e9, "k": 1e3, "thousand": 1e3}
FY = re.compile(r"FY\s?(\d{4})")
NUM = re.compile(r"\(?-?\d[\d,]*(?:\.\d+)?\)?%?")
AMOUNT = re.compile(r"\$?\s?(\d[\d,]*(?:\.\d+)?)\s*(%|million|billion|thousand|mm|bn|m|b|k)?(?![A-Za-z])", re.I)


@dataclass
class Fact:
    metric: str
    period: str
    value: float        # dollars for money, percentage points for pct
    unit: str           # usd / pct
    doc: str            # deck / audited / model
    file: str
    loc: str
    text: str = ""
    precision: float = 0.0
    derived: bool = False
    overridden: bool = False


def clean_label(s):
    s = re.sub(r"\(.*?\)", "", str(s)).replace("$", "").replace(":", "")
    return re.sub(r"\s+", " ", s).strip().lower()


def unit_hint(label):
    """scale or '%' from a label like 'Revenue ($M)' or 'EBITDA margin (%)'."""
    m = re.search(r"\((.*?)\)", str(label))
    if not m:
        return None
    h = m.group(1).lower().replace(" ", "")
    if "%" in h:
        return "%"
    if h in ("$m", "$mm", "usdm", "$millions"):
        return 1e6
    if h in ("$k", "$000", "$000s", "$thousands"):
        return 1e3
    return None


def parse_cell(s):
    """returns (value, decimals, is_pct) or None."""
    s = str(s).strip()
    m = NUM.fullmatch(s.replace("$", "").replace(" ", ""))
    if not m:
        return None
    raw = m.group(0)
    neg = raw.startswith("(") or raw.startswith("-")
    pct = raw.endswith("%")
    digits = raw.strip("()-%").replace(",", "")
    dec = len(digits.split(".")[1]) if "." in digits else 0
    v = float(digits)
    return (-v if neg else v), dec, pct


def make_fact(metric, period, num, dec, scale_or_pct, doc, file, loc, text):
    pct = metric in PCT_METRICS or scale_or_pct == "%"
    if pct:
        value, unit, prec = num, "pct", 0.5 * 10 ** -dec
    else:
        scale = scale_or_pct if isinstance(scale_or_pct, (int, float)) else 1.0
        value, unit, prec = num * scale, "usd", 0.5 * 10 ** -dec * scale
    if metric == "capex":
        value = abs(value)
    return Fact(metric, period, value, unit, doc, file, loc, text.strip()[:160], prec)


def scan_text(line, doc, file, loc, default_scale=1.0):
    """'Net debt (FY2024): $18.0M' style lines and 'x% of revenue' customer notes."""
    out = []
    fy = FY.findall(line)
    low = line.lower()
    if "customer" in low and "%" in line and "revenue" in low and fy:
        m = re.search(r"(\d+(?:\.\d+)?)\s*%", line)
        dec = len(m.group(1).split(".")[1]) if "." in m.group(1) else 0
        out.append(make_fact("customer_concentration", "FY" + fy[0], float(m.group(1)), dec, "%", doc, file, loc, line))
        return out
    m = re.match(r"^\s*[•\-\*]?\s*([A-Za-z][A-Za-z &/]*?)\s*(?:\((FY\s?\d{4})\))?\s*:\s*(.+)$", line)
    if not m:
        return out
    metric = LABEL_TO_METRIC.get(clean_label(m.group(1)))
    if not metric:
        return out
    period = m.group(2) or (("FY" + fy[0]) if fy else None)
    a = AMOUNT.search(m.group(3))
    if not period or not a:
        return out
    num_s, unit = a.group(1).replace(",", ""), (a.group(2) or "").lower()
    dec = len(num_s.split(".")[1]) if "." in num_s else 0
    hint = "%" if unit == "%" else SCALE.get(unit, default_scale)
    period = re.sub(r"\s", "", period)
    out.append(make_fact(metric, period, float(num_s), dec, hint, doc, file, loc, line))
    return out


def read_pdf(path, doc):
    import pdfplumber
    facts = []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            lines = (page.extract_text() or "").splitlines()
            periods, scale = [], 1.0
            for ln in lines:
                f = FY.findall(ln)
                if len(f) >= 2 and not re.search(r"[a-z]{4}", ln.replace("FY", "")):
                    periods = ["FY" + y for y in f]
                m = re.search(r"in (?:us )?(thousands|millions)", ln, re.I)
                if m:
                    scale = 1e3 if m.group(1).lower() == "thousands" else 1e6
            for ln in lines:
                facts += scan_text(ln, doc, path.name, f"p.{pno}")
                m = re.match(r"^(.*?[A-Za-z\)])\s+((?:\(?-?\d[\d,]*(?:\.\d+)?\)?%?\s*)+)$", ln.strip())
                if not m or not periods:
                    continue
                label, nums = clean_label(m.group(1)), NUM.findall(m.group(2))
                if label.startswith("segment ") or label == "total segment revenue":
                    metric = "segment_revenue" if label.startswith("segment ") else "segment_total"
                else:
                    metric = LABEL_TO_METRIC.get(label)
                if not metric:
                    continue
                for per, n in zip(periods, nums):
                    c = parse_cell(n)
                    if c:
                        facts.append(make_fact(metric, per, c[0], c[1], "%" if c[2] else scale, doc, path.name, f"p.{pno}", ln))
    return facts


def read_pptx(path, doc):
    from pptx import Presentation
    facts = []
    for sno, slide in enumerate(Presentation(str(path)).slides, 1):
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                periods = ["FY" + y for y in FY.findall(" ".join(rows[0]))]
                for r in rows[1:]:
                    metric = LABEL_TO_METRIC.get(clean_label(r[0]))
                    hint = unit_hint(r[0])
                    if not metric:
                        continue
                    for per, cell in zip(periods, r[1:]):
                        c = parse_cell(cell)
                        if c:
                            facts.append(make_fact(metric, per, c[0], c[1], hint or 1.0, doc, path.name,
                                                   f"slide {sno}", " | ".join(r)))
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in para.runs)
                    facts += scan_text(txt, doc, path.name, f"slide {sno}")
    return facts


def read_xlsx(path, doc):
    import openpyxl
    facts = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        scale, header_row, cols = 1.0, None, {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6)):
            for c in row:
                if isinstance(c.value, str):
                    t = c.value.lower()
                    if "thousand" in t or "000" in t:
                        scale = 1e3
                    elif "million" in t:
                        scale = 1e6
        for row in ws.iter_rows():
            hits = {c.column: "FY" + FY.search(str(c.value)).group(1) for c in row if c.value and FY.search(str(c.value))}
            if len(hits) >= 2 and header_row is None:
                header_row, cols = row[0].row, hits
                continue
            if header_row is None or row[0].row <= header_row:
                continue
            label = next((c.value for c in row if isinstance(c.value, str)), None)
            metric = LABEL_TO_METRIC.get(clean_label(label)) if label else None
            if not metric:
                continue
            for c in row:
                if c.column in cols and isinstance(c.value, (int, float)):
                    dec = len(str(c.value).split(".")[1]) if "." in str(c.value) else 0
                    facts.append(make_fact(metric, cols[c.column], float(c.value), dec, scale, doc, path.name,
                                           f"{ws.title}!{c.coordinate}", f"{label}: {c.value}"))
    return facts


READERS = {".pptx": read_pptx, ".pdf": read_pdf, ".xlsx": read_xlsx}


def read_deal_room(folder):
    facts, skipped = [], []
    for p in sorted(Path(folder).iterdir()):
        ext = p.suffix.lower()
        if ext in READERS:
            facts += READERS[ext](p, ROLE_BY_EXT[ext])
        elif p.is_file() and p.name != "truth.json" and not p.name.endswith(".py"):
            skipped.append(p.name)
    return facts, skipped
