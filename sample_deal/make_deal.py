"""
Makes a fake Indian deal room for an invented company: a management deck, audited
Ind AS statements and an operating model, with a few planted inconsistencies and
a truth.json saying exactly what was planted. Nothing here is real company data.

    python sample_deal/make_deal.py sample_deal --seed 7
"""
import argparse
import json
import random
import textwrap
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

NAME = "Nilgiri Precision Components Limited"
TAX = 0.2517
ISSUES = {
    "deck_revenue": ("mismatch", "revenue", "FY2024"),
    "deck_ebitda": ("mismatch", "ebitda", "FY2024"),
    "deck_pat": ("mismatch", "net_income", "FY2024"),
    "deck_net_debt": ("mismatch", "net_debt", "FY2024"),
    "deck_gross_margin": ("mismatch", "gross_margin", "FY2024"),
    "model_capex": ("mismatch", "capex", "FY2024"),
    "model_revenue_py": ("mismatch", "revenue", "FY2023"),
    "customer_conc": ("mismatch", "customer_concentration", "FY2024"),
    "orderbook_unsupported": ("unsupported", "order_book", "FY2024"),
    "segment_sum": ("internal", "revenue", "FY2024"),
    "deck_margin_math": ("internal", "ebitda_margin", "FY2024"),
    # footnote cases
    "fn_ebitda_full": ("mismatch", "ebitda", "FY2024"),
    "fn_ebitda_partial": ("mismatch", "ebitda", "FY2024"),
    "fn_ebitda_noamt": ("mismatch", "ebitda", "FY2024"),
    "fn_netdebt_lease": ("mismatch", "net_debt", "FY2024"),
    "fn_audited_lease": ("mismatch", "net_debt", "FY2024"),
    "fn_capex_comment": ("mismatch", "capex", "FY2024"),
    "fn_contingent": ("disclosure", "contingent_liabilities", "FY2024"),
    "fn_emphasis": ("disclosure", "audit_emphasis", "FY2024"),
    "fn_going_concern": ("disclosure", "going_concern", "FY2024"),
}
# what the footnote check should say for the cases above
EXPLAINED = {"fn_ebitda_full": "full", "fn_ebitda_partial": "partial", "fn_ebitda_noamt": "noted",
             "fn_netdebt_lease": "full", "fn_audited_lease": "full", "fn_capex_comment": "full"}
GROUPS = [{"deck_ebitda", "fn_ebitda_full", "fn_ebitda_partial", "fn_ebitda_noamt"},
          {"deck_net_debt", "fn_netdebt_lease", "fn_audited_lease"},
          {"model_capex", "fn_capex_comment"}]
MARKS = ["(1)", "*", "¹", "†"]


def inr(x, dec=2):
    """Indian digit grouping: 12,34,567.89"""
    s = f"{abs(x):.{dec}f}"
    ip, _, fp = s.partition(".")
    if len(ip) > 3:
        head, tail, parts = ip[:-3], ip[-3:], []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        ip = ",".join(([head] if head else []) + parts + [tail])
    out = ip + ("." + fp if fp else "")
    return f"({out})" if x < 0 else out


def company(rng):
    """three years (FY2022..FY2024), all in rupees crore, two decimals."""
    r0 = rng.uniform(350, 900)
    g = [rng.uniform(0.06, 0.16) for _ in range(2)]
    rev = [r0, r0 * (1 + g[0]), r0 * (1 + g[0]) * (1 + g[1])]
    d = {k: [] for k in ("rev", "oi", "material", "employee", "other", "fin", "da", "total_exp", "ebitda",
                         "pbt", "tax", "pat", "capex", "wc")}
    gm, em = rng.uniform(0.34, 0.44), rng.uniform(0.13, 0.22)
    for r in rev:
        r = round(r, 2)
        mat = round(r * (1 - gm + rng.uniform(-.01, .01)), 2)
        opex = r * (1 - em + rng.uniform(-.01, .01))
        emp = round((opex - mat) * rng.uniform(.38, .5), 2)
        oth = round(opex - mat - emp, 2)
        da = round(r * rng.uniform(.03, .045), 2)
        fin = round(r * rng.uniform(.012, .03), 2)
        oi = round(r * rng.uniform(.004, .015), 2)
        te = round(mat + emp + oth + fin + da, 2)
        pbt = round(r + oi - te, 2)
        tax = round(pbt * TAX, 2)
        for k, v in (("rev", r), ("oi", oi), ("material", mat), ("employee", emp), ("other", oth), ("fin", fin),
                     ("da", da), ("total_exp", te), ("ebitda", round(r - mat - emp - oth, 2)), ("pbt", pbt),
                     ("tax", tax), ("pat", round(pbt - tax, 2)), ("capex", round(r * rng.uniform(.04, .08), 2)),
                     ("wc", round(r * rng.uniform(-.02, .02), 2))):
            d[k].append(v)
    d["cash"] = [round(d["rev"][i] * rng.uniform(.03, .10), 2) for i in (1, 2)]
    d["debt_nc"] = [round(d["rev"][i] * rng.uniform(.12, .28), 2) for i in (1, 2)]
    d["debt_c"] = [round(d["rev"][i] * rng.uniform(.06, .15), 2) for i in (1, 2)]
    sh = [rng.uniform(.35, .5), rng.uniform(.25, .35)]
    sh.append(1 - sum(sh))
    d["shares"] = sh
    d["top_customer"] = round(rng.uniform(10, 24))
    return d


def font(c):
    """a font with the rupee sign; falls back to Helvetica and 'Rs.'"""
    try:
        import matplotlib
        p = Path(matplotlib.__file__).parent / "mpl-data/fonts/ttf/DejaVuSans.ttf"
        pdfmetrics.registerFont(TTFont("DV", str(p)))
        return "DV", "₹"
    except Exception:
        return "Helvetica", "Rs."


def build(out_dir, seed=7, issues=None, hard=False):
    # hard: other label wordings, smaller gaps, projection columns and a whole-crore
    # rounded figure that is correct and should not be flagged
    rng = random.Random(seed)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    d = company(rng)
    if issues is None:
        issues = [k for k in ISSUES if rng.random() < 0.3]
        while len(issues) < 2:
            issues = sorted(set(issues + [rng.choice(list(ISSUES))]))
        for g in GROUPS:                       # one problem per figure
            kept = [k for k in issues if k in g]
            issues = [k for k in issues if k not in g] + kept[:1]
    has = set(issues).__contains__
    pdf_unit, model_unit = rng.choice(["crore", "lakh"]), rng.choice(["crore", "lakh"])
    pdf_k, model_k = (1 if pdf_unit == "crore" else 100), (1 if model_unit == "crore" else 100)
    pdf_style = rng.choice(["fy", "march"])

    rev, ebitda, pat, capex = d["rev"], d["ebitda"], d["pat"], d["capex"]
    lease = None
    if has("fn_netdebt_lease") or has("fn_audited_lease"):
        lease = round(rev[2] * rng.uniform(.015, .04), 1)
    if has("fn_audited_lease"):                # audited borrowings include leases; the deck leaves them out
        d["debt_nc"][1] = round(d["debt_nc"][1] + lease, 2)
    aud_nd = d["debt_nc"][1] + d["debt_c"][1] - d["cash"][1]
    aud_gm = [100 * (r - m) / r for r, m in zip(rev, d["material"])]

    # ---- management deck, in rupees crore, one decimal
    deck_rev = [round(x, 1) for x in rev]
    deck_ebitda = [round(x, 1) for x in ebitda]
    deck_pat = [round(x, 1) for x in pat]
    deck_gm = [round(x, 1) for x in aud_gm]
    deck_nd = round(aud_nd, 1)
    deck_top = d["top_customer"]
    rev_lab = rng.choice(["Revenue from operations", "Net sales", "Total revenue"]) if hard else "Revenue"
    pat_lab = rng.choice(["PAT", "Profit after tax", "Net profit"]) if hard else "PAT"
    lo = (lambda a, b: (a, b)) if not hard else (lambda a, b: (a / 3.5, b / 4))
    if has("deck_revenue"):
        a, b = lo(.04, .15)
        deck_rev[2] = round(rev[2] * (1 + rng.uniform(a, b)), 1)
    if has("deck_ebitda"):
        a, b = lo(.05, .20)
        deck_ebitda[2] = round(ebitda[2] * (1 + rng.uniform(a, b)), 1)
    if has("deck_pat"):
        a, b = lo(.06, .18)
        deck_pat[2] = round(pat[2] * (1 + rng.uniform(a, b)), 1)
    if has("deck_net_debt"):
        deck_nd = round(aud_nd * (1 - rng.uniform(.25, .6)), 1)
    if has("deck_gross_margin"):
        deck_gm[2] = round(aud_gm[2] + (rng.uniform(.8, 2.0) if hard else rng.uniform(2.5, 6)), 1)
    if has("customer_conc"):
        deck_top = max(d["top_customer"] - round(rng.uniform(6, 12)), 2)
    mk1, mk2 = rng.sample(MARKS, 2)
    ebitda_note = None
    if has("fn_ebitda_full") or has("fn_ebitda_partial") or has("fn_ebitda_noamt"):
        e = ebitda[2]
        if has("fn_ebitda_full"):
            x = round(e * rng.uniform(.04, .12), 1)
            gap = x
            ebitda_note = rng.choice([f"Adjusted EBITDA; excludes exceptional items of ₹{x:.1f} Cr.",
                                      f"Before one-time restructuring costs of ₹{x:.1f} Cr.",
                                      f"Excludes ESOP expense of ₹{x:.1f} Cr."])
        elif has("fn_ebitda_partial"):
            x = round(e * rng.uniform(.03, .06), 1)
            gap = x + round(e * rng.uniform(.02, .04), 1)
            ebitda_note = f"Adjusted EBITDA; excludes exceptional items of ₹{x:.1f} Cr."
        else:
            gap = round(e * rng.uniform(.05, .15), 1)
            ebitda_note = "Adjusted EBITDA as per management definition."
        deck_ebitda[2] = round(e + gap, 1)
    nd_note = None
    if has("fn_netdebt_lease"):
        deck_nd = round(aud_nd + lease, 1)
        nd_note = f"Includes lease liabilities of ₹{lease:.1f} Cr."
    if has("fn_audited_lease"):
        deck_nd = round(aud_nd - lease, 1)
    deck_margin = [round(100 * e / r, 1) for e, r in zip(deck_ebitda, deck_rev)]
    if has("deck_margin_math"):
        deck_margin[2] = round(deck_margin[2] + rng.uniform(2, 5), 1)

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = NAME
    s.placeholders[1].text = "Management presentation (fictional company, synthetic data)"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Financial highlights"
    heads = ["Particulars (₹ Cr)", "FY23", "FY24"] + (["FY25E"] if hard else [])
    proj = [round(rev[2] * 1.13, 1), round(ebitda[2] * 1.18, 1), round(pat[2] * 1.2, 1)]
    rows = [heads,
            [rev_lab, f"{deck_rev[1]:.1f}", f"{deck_rev[2]:.1f}"] + ([f"{proj[0]:.1f}"] if hard else []),
            [f"EBITDA {mk1}" if ebitda_note else "EBITDA", f"{deck_ebitda[1]:.1f}", f"{deck_ebitda[2]:.1f}"] + ([f"{proj[1]:.1f}"] if hard else []),
            ["EBITDA margin (%)", f"{deck_margin[1]:.1f}", f"{deck_margin[2]:.1f}"] + ([f"{100 * proj[1] / proj[0]:.1f}"] if hard else []),
            ["Gross margin (%)", f"{deck_gm[1]:.1f}", f"{deck_gm[2]:.1f}"] + ([f"{deck_gm[2]:.1f}"] if hard else []),
            [pat_lab, f"{deck_pat[1]:.1f}", f"{deck_pat[2]:.1f}"] + ([f"{proj[2]:.1f}"] if hard else [])]
    t = s.shapes.add_table(len(rows), len(heads), Inches(0.7), Inches(1.8), Inches(8.6), Inches(2.6)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v
    fbox = s.shapes.add_textbox(Inches(0.7), Inches(4.7), Inches(8.6), Inches(1)).text_frame
    flines = [f"{mk1} {ebitda_note}"] if ebitda_note else []
    flines.append("Source: Audited financial statements; figures rounded to one decimal.")
    if hard:
        flines.append(f"Note: {rev_lab} includes other operating income of ₹{round(rev[2] * rng.uniform(.002, .006), 1)} Cr.")
    fbox.text = flines[0]
    for ln in flines[1:]:
        fbox.add_paragraph().text = ln
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Balance sheet, customers and order book"
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(8.6), Inches(3)).text_frame
    bullets = [f"Net debt (FY24): ₹{deck_nd:.1f} Cr" + (f" {mk2}" if nd_note else ""),
               f"Cash (FY24): ₹{d['cash'][1]:.1f} Cr",
               f"Top customer: {deck_top}% of FY24 revenue"]
    if hard:
        bullets.append(f"{rev_lab} (FY23): ₹{round(rev[1])} Cr")
    if has("orderbook_unsupported"):
        bullets.append(f"Order book (FY24): ₹{rev[2] * rng.uniform(.5, .9):,.1f} Cr")
    box.text = bullets[0]
    for b in bullets[1:]:
        box.add_paragraph().text = b
    if nd_note:
        nbox = s.shapes.add_textbox(Inches(0.7), Inches(4.9), Inches(8.6), Inches(0.5)).text_frame
        nbox.text = f"{mk2} {nd_note}"
    prs.save(out / "management_deck.pptx")

    # ---- audited statements (two years, Schedule III style, note-number column)
    segs = [[round(rev[i] * sh, 2) for i in (1, 2)] for sh in d["shares"]]
    for i in range(2):
        segs[2][i] = round(rev[i + 1] - segs[0][i] - segs[1][i], 2)
    if has("segment_sum"):
        k = rng.uniform(.02, .06)
        segs[2][1] = round(segs[2][1] - rev[2] * k, 2)
    fname, sym = font(None)
    cont = round(ebitda[2] * rng.uniform(.1, .6), 2)
    c = canvas.Canvas(str(out / "audited_financials.pdf"), pagesize=A4)
    unit_line = f"({sym} in {pdf_unit}{'s' if pdf_unit == 'lakh' and rng.random() < .5 else ''})"
    cols = ("FY2023-24", "FY2022-23") if pdf_style == "fy" else ("31 March 2024", "31 March 2023")
    bs_cols = ("31 March 2024", "31 March 2023")

    def page(title, header, rows, extra=(), note_col=True):
        y = 790
        c.setFont(fname, 13); c.drawString(60, y, f"{NAME}"); y -= 18
        c.setFont(fname, 12); c.drawString(60, y, title); y -= 20
        c.setFont(fname, 9)
        if unit_line:
            c.drawString(60, y, unit_line); y -= 16
        if header:
            c.drawString(60, y, "Particulars")
            if note_col:
                c.drawRightString(320, y, "Note")
            for x, h in zip((430, 520), header):
                c.drawRightString(x, y, h)
            y -= 18
        c.setFont(fname, 10)
        for label, note, vals in rows:
            c.drawString(60, y, label)
            if note and note_col:
                c.drawRightString(320, y, str(note))
            for x, v in zip((430, 520), vals):
                c.drawRightString(x, y, inr(v * pdf_k))
            y -= 16
        for line in extra:
            c.drawString(60, y, line); y -= 14
        c.showPage()

    report_lines = ["In our opinion the Ind AS financial statements give the information required by the Companies Act, 2013",
                    "in the manner so required and give a true and fair view (fictional company, synthetic data).",
                    "", "Key Audit Matters", "Revenue recognition and the valuation of inventory were the matters of most significance.",
                    "", "Responsibilities of Management and Auditor",
                    "Management is responsible for assessing the company's ability to continue as a going concern.",
                    "We conclude on the appropriateness of the going concern basis and whether a material uncertainty exists."]
    if has("fn_emphasis"):
        report_lines += ["", "Emphasis of Matter",
                         *textwrap.wrap("We draw attention to Note 41, which describes the uncertainty relating to the outcome of a "
                                        "tax assessment. Our opinion is not modified in respect of this matter.", 100)]
    if has("fn_going_concern"):
        report_lines += ["", "Material Uncertainty Related to Going Concern",
                         *textwrap.wrap("We draw attention to Note 42, which indicates that the company's current liabilities exceed "
                                        "its current assets and its lenders have not yet agreed to a revised repayment schedule.", 100)]
    page("Independent Auditor's Report", [], [], report_lines)
    last = lambda k: [d[k][2], d[k][1]]
    page("Balance Sheet as at 31 March 2024", bs_cols,
         [("Cash and cash equivalents", 9, [d["cash"][1], d["cash"][0]]),
          ("Other assets", 10, [round(rev[2] * .9, 2), round(rev[1] * .9, 2)]),
          ("Non-current borrowings" + (" (1)" if has("fn_audited_lease") else ""), 17, [d["debt_nc"][1], d["debt_nc"][0]]),
          ("Current borrowings", 21, [d["debt_c"][1], d["debt_c"][0]]),
          ("Other liabilities", 22, [round(rev[2] * .2, 2), round(rev[1] * .2, 2)])],
         textwrap.wrap(f"(1) Includes lease liabilities of {sym}{lease * pdf_k:,.2f} {pdf_unit} recognised under Ind AS 116 "
                       "and presented within borrowings.", 62) if has("fn_audited_lease") else ())
    page("Statement of Profit and Loss for the year ended 31 March 2024", cols,
         [("Revenue from operations", 23, last("rev")), ("Other income", 24, last("oi")),
          ("Total income", "", [round(a + b, 2) for a, b in zip(last("rev"), last("oi"))]),
          ("Cost of materials consumed", 25, last("material")), ("Employee benefits expense", 26, last("employee")),
          ("Finance costs", 27, last("fin")), ("Depreciation and amortisation expense", 28, last("da")),
          ("Other expenses", 29, last("other")), ("Total expenses", "", last("total_exp")),
          ("Profit before tax", "", last("pbt")), ("Tax expense", 30, last("tax")),
          ("Profit for the year", "", last("pat"))], ["Note: Figures in brackets are negative."])
    page("Statement of Cash Flows for the year ended 31 March 2024", cols,
         [("Profit before tax", "", last("pbt")), ("Depreciation and amortisation expense", "", last("da")),
          ("Change in working capital", "", last("wc")),
          ("Purchase of property, plant and equipment", "", [-d["capex"][2], -d["capex"][1]])], note_col=False)
    page("Notes to the financial statements", cols,
         [(f"Segment {n}", "", [segs[i][1], segs[i][0]]) for i, n in enumerate(["Auto components", "Industrial", "Exports"])]
         + [("Total segment revenue", "", [round(sum(sg[1] for sg in segs), 2), round(sum(sg[0] for sg in segs), 2)])],
         ["Note 38. Segment information",
          f"Note 39. One customer contributed {d['top_customer']}% of revenue from operations in FY2024."]
         + ([f"Note 40. Contingent liabilities not provided for amounted to {sym}{cont * pdf_k:,.2f} {pdf_unit} as at 31 March 2024."]
            if has("fn_contingent") else []), note_col=False)
    c.save()

    # ---- operating model
    def jitter(x):
        return x * (1 + rng.uniform(-4e-4, 4e-4))

    m_rev = [jitter(x) for x in rev]
    m_capex = [jitter(x) for x in capex]
    if has("model_revenue_py"):
        a, b = (.012, .03) if hard else (.03, .08)
        m_rev[1] = rev[1] * (1 + rng.choice([-1, 1]) * rng.uniform(a, b))
    if has("model_capex"):
        a, b = (.03, .08) if hard else (.2, .4)
        m_capex[2] = capex[2] * (1 + rng.choice([-1, 1]) * rng.uniform(a, b))
    cap_x = round(capex[2] * rng.uniform(.05, .15), 2)
    if has("fn_capex_comment"):
        m_capex[2] = capex[2] + cap_x
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Operating model, INR {model_unit}"
    ws.append([])
    heads = ["Line item", "FY22", "FY23", "FY24"] + (["FY25E"] if hard else [])
    ws.append(heads)
    m_rev_lab = rng.choice(["Revenue", "Revenue from operations", "Net sales"]) if hard else "Revenue"
    m_cap_lab = rng.choice(["Capex", "Capital expenditure"]) if hard else "Capex"
    m_pat_lab = rng.choice(["PAT", "Profit after tax"]) if hard else "PAT"
    cv = lambda xs, extra: [round(x * model_k, 1 if model_unit == "lakh" else 2) for x in xs] + \
        ([round(extra * model_k, 1 if model_unit == "lakh" else 2)] if hard else [])
    for label, vals, ext in ((m_rev_lab, m_rev, rev[2] * 1.12), ("EBITDA", [jitter(x) for x in ebitda], ebitda[2] * 1.17),
                             (m_pat_lab, [jitter(x) for x in pat], pat[2] * 1.2), (m_cap_lab, m_capex, capex[2] * 1.1)):
        ws.append([label, *cv(vals, ext)])
    if has("fn_capex_comment"):
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == m_cap_lab:
                txt = (f"Includes capitalised borrowing costs of ₹{cap_x:.2f} Cr." if model_unit == "crore"
                       else f"Includes capitalised borrowing costs of ₹{cap_x * 100:.1f} lakh.")
                ws.cell(row=r, column=4).comment = Comment(txt, "model owner")
    w2 = wb.create_sheet("Debt")
    w2["A1"] = f"INR {model_unit}"
    w2.append([])
    w2.append(["Line item", "FY23", "FY24"])
    w2.append(["Cash", *cv([jitter(d["cash"][0]), jitter(d["cash"][1])], 0)[:2]])
    w2.append(["Total debt", *cv([jitter(d["debt_nc"][0] + d["debt_c"][0]), jitter(d["debt_nc"][1] + d["debt_c"][1])], 0)[:2]])
    wb.save(out / "operating_model.xlsx")

    truth = [{"issue": k, "kind": ISSUES[k][0], "metric": ISSUES[k][1], "period": ISSUES[k][2],
              "explained": EXPLAINED.get(k, "")} for k in issues]
    (out / "truth.json").write_text(json.dumps({"seed": seed, "hard": hard, "pdf_unit": pdf_unit,
                                                "model_unit": model_unit, "planted": truth}, indent=2))
    return truth


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--hard", action="store_true")
    ap.add_argument("--issues", help="comma separated issue names; default is a random pick")
    a = ap.parse_args()
    picked = a.issues.split(",") if a.issues else None
    for t in build(a.out, a.seed, issues=picked, hard=a.hard):
        print("planted:", t["issue"])
